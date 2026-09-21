import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from utils.constants import Decree30Rules, ErrorMessages, HighlightColors

class RuleEngine:
    def __init__(self, file_path):
        self.file_path = file_path
        self.doc = Document(file_path)
        self.error_log = []
        self.in_noi_nhan_section = False 

    def check_and_format(self, add_note=True):
        self.in_noi_nhan_section = False
        
        for i, para in enumerate(self.doc.paragraphs):
            self._process_paragraph(para, f"Đoạn {i+1}", add_note)
            
        for t_idx, table in enumerate(self.doc.tables):
            for r_idx, row in enumerate(table.rows):
                for c_idx, cell in enumerate(row.cells):
                    self.in_noi_nhan_section = False 
                    for p_idx, para in enumerate(cell.paragraphs):
                        self._process_paragraph(para, f"Bảng {t_idx+1} - Ô({r_idx+1},{c_idx+1})", add_note)

        return self.error_log

    def _log_error(self, vi_tri, text, err_msg, fix_guide, para, add_note):
        self.error_log.append({
            "vi_tri": vi_tri,
            "noi_dung": (text[:40] + "...") if len(text) > 40 else text,
            "loi": err_msg,
            "cach_sua": fix_guide
        })
        if add_note:
            note_run = para.add_run(f" {err_msg} - {fix_guide}")
            note_run.bold = True
            note_run.font.highlight_color = HighlightColors.NOTE_COLOR

    def _process_paragraph(self, para, vi_tri, add_note):
        for run in para.runs:
            if run.text and ("[LỖI:" in run.text or "Thao tác:" in run.text):
                run.text = "" 
            if run.font.highlight_color in [HighlightColors.FORMAT_ERROR, HighlightColors.NOTE_COLOR]:
                run.font.highlight_color = None 

        text = para.text.strip()
        if not text:
            return 
            
        original_runs = list(para.runs)

        is_quoc_hieu = "CỘNG HÒA" in text.upper() or "CỘNG HOÀ" in text.upper()
        is_tieu_ngu = "Độc lập" in text

        # ==========================================
        # A. KIỂM TRA QUỐC HIỆU & TIÊU NGỮ
        # ==========================================
        if is_quoc_hieu:
            if text not in ["CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", "CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM"]:
                self._log_error(vi_tri, text, ErrorMessages.ERR_QUOC_HIEU_TEXT, "Sửa thành: CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", para, add_note)
            
            for run in original_runs:
                if run.text.strip():
                    is_bold = run.bold if run.bold is not None else False
                    if not is_bold:
                        self._log_error(vi_tri, text, "[LỖI: Quốc hiệu thiếu in ĐẬM]", "Bôi đen -> Bấm Ctrl+B", para, add_note)
                        run.font.highlight_color = HighlightColors.FORMAT_ERROR
                    
                    if run.font.size and run.font.size.pt not in [12.0, 12.5, 13.0]:
                        self._log_error(vi_tri, text, f"[LỖI: Cỡ chữ Quốc hiệu đang là {run.font.size.pt}, yêu cầu 12-13]", "Chỉnh lại Size 12 hoặc 13", para, add_note)
                        run.font.highlight_color = HighlightColors.FORMAT_ERROR
            return 

        elif is_tieu_ngu:
            if text != "Độc lập - Tự do - Hạnh phúc":
                self._log_error(vi_tri, text, ErrorMessages.ERR_TIEU_NGU_TEXT, "Dùng dấu gạch nối '-', khoảng trắng 2 bên", para, add_note)
            
            for run in original_runs:
                if run.text.strip():
                    is_bold = run.bold if run.bold is not None else False
                    if not is_bold:
                        self._log_error(vi_tri, text, "[LỖI: Tiêu ngữ thiếu in ĐẬM]", "Bôi đen -> Bấm Ctrl+B", para, add_note)
                        run.font.highlight_color = HighlightColors.FORMAT_ERROR
                    
                    if run.font.size and (run.font.size.pt < 13.0 or run.font.size.pt > 14.0):
                        self._log_error(vi_tri, text, f"[LỖI: Cỡ chữ Tiêu ngữ đang là {run.font.size.pt}, yêu cầu 13-14]", "Chỉnh lại Size 13 hoặc 14", para, add_note)
                        run.font.highlight_color = HighlightColors.FORMAT_ERROR
            return 

        # ==========================================
        # B. KIỂM TRA NƠI NHẬN
        # ==========================================
        if text.lower().startswith("nơi nhận"):
            self.in_noi_nhan_section = True
            
            if para.alignment not in [WD_ALIGN_PARAGRAPH.LEFT, None]:
                self._log_error(vi_tri, text, ErrorMessages.ERR_NOI_NHAN_ALIGN, "Bấm Ctrl+L (Align Left)", para, add_note)
            
            has_font_error = False
            for run in original_runs:
                if run.text.strip():
                    is_bold = run.bold if run.bold is not None else False
                    is_italic = run.italic if run.italic is not None else False
                    
                    if not is_bold or not is_italic:
                        has_font_error = True
                        run.font.highlight_color = HighlightColors.FORMAT_ERROR
                        
                    if run.font.size and run.font.size.pt != 12.0:
                        has_font_error = True
                        run.font.highlight_color = HighlightColors.FORMAT_ERROR
                        
            if has_font_error:
                self._log_error(vi_tri, text, ErrorMessages.ERR_NOI_NHAN_FONT, "Bôi đen -> Ctrl+B, Ctrl+I, Size 12", para, add_note)
            return

        if self.in_noi_nhan_section:
            if text.startswith("-"):
                if para.alignment not in [WD_ALIGN_PARAGRAPH.LEFT, None]:
                    self._log_error(vi_tri, text, ErrorMessages.ERR_NOI_NHAN_ALIGN, "Bấm Ctrl+L (Align Left)", para, add_note)
                
                has_font_error = False
                for run in original_runs:
                    if run.text.strip():
                        is_bold = run.bold if run.bold is not None else False
                        is_italic = run.italic if run.italic is not None else False
                        
                        if is_bold or is_italic:
                            has_font_error = True
                            run.font.highlight_color = HighlightColors.FORMAT_ERROR
                            
                        if run.font.size and run.font.size.pt != 11.0:
                            has_font_error = True
                            run.font.highlight_color = HighlightColors.FORMAT_ERROR
                            
                if has_font_error:
                    self._log_error(vi_tri, text, ErrorMessages.ERR_NOI_NHAN_LIST, "Bôi đen -> Bỏ Đậm/Nghiêng, Size 11", para, add_note)
                return
            else:
                self.in_noi_nhan_section = False 

        # ==========================================
        # C. KIỂM TRA VĂN BẢN CHUNG
        # ==========================================
        para_errors = []
        
        for run in original_runs:
            if run.text.strip():
                if run.font.name != Decree30Rules.FONT_NAME and run.font.name is not None:
                    para_errors.append(ErrorMessages.ERR_FONT)
                    run.font.highlight_color = HighlightColors.FORMAT_ERROR
                
                if run.font.size:
                    size_pt = run.font.size.pt
                    if size_pt < Decree30Rules.FONT_SIZE_MIN or size_pt > Decree30Rules.FONT_SIZE_MAX:
                        para_errors.append(ErrorMessages.ERR_SIZE)
                        run.font.highlight_color = HighlightColors.FORMAT_ERROR
                        run.font.size = Pt(14) 

        is_title = text.isupper() and len(text) < 100 and not is_quoc_hieu
        is_abstract = text.startswith("Về việc") or text.startswith("DỰ TOÁN") or (text.startswith("BÁO CÁO") and not is_title)
        is_attachment = text.startswith("(Đính kèm") or text.startswith("(Kèm theo")
        is_recipient = text.startswith("Kính gửi:")
        is_date_line = ", ngày" in text.lower() and "năm" in text.lower()

        # Check Canh giữa
        if is_title or is_abstract or is_attachment or is_recipient:
            if para.alignment != WD_ALIGN_PARAGRAPH.CENTER:
                err_msg = ErrorMessages.ERR_ALIGN_CENTER
                fix_guide = "Thao tác: Bôi đen dòng này -> Bấm Ctrl+E (Center)"
                self._log_error(vi_tri, text, err_msg, fix_guide, para, add_note)
                if add_note and original_runs:
                    original_runs[0].font.highlight_color = HighlightColors.FORMAT_ERROR
            
            # KIỂM TRA CHUYÊN SÂU CHO TÊN LOẠI VÀ TRÍCH YẾU
            if is_title:
                has_bold_err = False
                for run in original_runs:
                    if run.text.strip():
                        is_bold = run.bold if run.bold is not None else False
                        if not is_bold:
                            has_bold_err = True
                            run.font.highlight_color = HighlightColors.FORMAT_ERROR
                if has_bold_err:
                    self._log_error(vi_tri, text, ErrorMessages.ERR_TITLE_FONT, "Bôi đen -> Bấm Ctrl+B", para, add_note)
                    
            if is_abstract:
                has_bold_err = False
                has_underline_err = False
                for run in original_runs:
                    if run.text.strip():
                        is_bold = run.bold if run.bold is not None else False
                        if not is_bold:
                            has_bold_err = True
                            run.font.highlight_color = HighlightColors.FORMAT_ERROR
                        if run.underline:
                            has_underline_err = True
                            run.font.highlight_color = HighlightColors.FORMAT_ERROR
                
                if has_bold_err:
                    self._log_error(vi_tri, text, ErrorMessages.ERR_ABSTRACT_FONT, "Bôi đen -> Bấm Ctrl+B", para, add_note)
                if has_underline_err:
                    self._log_error(vi_tri, text, ErrorMessages.ERR_ABSTRACT_LINE, "Tắt gạch chân (Ctrl+U) -> Insert > Shapes > Line", para, add_note)
                    
        elif is_date_line:
            has_font_error = False
            for run in original_runs:
                if run.text.strip():
                    is_italic = run.italic if run.italic is not None else False
                    if not is_italic:
                        has_font_error = True
                        run.font.highlight_color = HighlightColors.FORMAT_ERROR
                    
                    if run.font.size and (run.font.size.pt < 13.0 or run.font.size.pt > 14.0):
                        has_font_error = True
                        run.font.highlight_color = HighlightColors.FORMAT_ERROR
                        
            if has_font_error:
                self._log_error(vi_tri, text, ErrorMessages.ERR_DATE_FONT, "Bôi đen -> Bấm Ctrl+I, Size 13 hoặc 14", para, add_note)
                if add_note and original_runs:
                    original_runs[0].font.highlight_color = HighlightColors.FORMAT_ERROR
                    
        else:
            if len(text) > 20:
                if para.alignment != WD_ALIGN_PARAGRAPH.JUSTIFY:
                    err_msg = ErrorMessages.ERR_ALIGN
                    fix_guide = "Thao tác: Bôi đen đoạn này -> Bấm Ctrl+J (Justify)"
                    self._log_error(vi_tri, text, err_msg, fix_guide, para, add_note)
                    if add_note and original_runs:
                        original_runs[0].font.highlight_color = HighlightColors.FORMAT_ERROR

                indent = para.paragraph_format.first_line_indent
                if indent is None or not (0.95 <= indent.cm <= 1.3):
                    err_msg = ErrorMessages.ERR_INDENT
                    fix_guide = "Thao tác: Paragraph -> Special: First line -> 1.27 cm"
                    self._log_error(vi_tri, text, err_msg, fix_guide, para, add_note)
                    if add_note and original_runs:
                        original_runs[0].font.highlight_color = HighlightColors.FORMAT_ERROR

                space_after = para.paragraph_format.space_after
                space_before = para.paragraph_format.space_before
                
                pt_after = space_after.pt if space_after else 0
                pt_before = space_before.pt if space_before else 0
                
                if (pt_after + pt_before) < 6.0 or (pt_after > 0 and pt_before > 0):
                    err_msg = ErrorMessages.ERR_SPACING
                    fix_guide = "Thao tác: Paragraph -> Spacing: Before -> 0 pt, After -> 6 pt"
                    self._log_error(vi_tri, text, err_msg, fix_guide, para, add_note)
                    if add_note and original_runs:
                        original_runs[0].font.highlight_color = HighlightColors.FORMAT_ERROR

        font_errors = [e for e in para_errors if e in (ErrorMessages.ERR_FONT, ErrorMessages.ERR_SIZE)]
        if font_errors:
            unique_fonts = list(dict.fromkeys(font_errors))
            error_text = " | ".join(unique_fonts)
            self._log_error(vi_tri, text, error_text, "Quét chọn -> Chỉnh Font/Size trên thanh công cụ", para, add_note)

    def save_output(self):
        dir_name = os.path.dirname(self.file_path)
        base_name = os.path.basename(self.file_path)
        name, ext = os.path.splitext(base_name)
        
        output_name = f"{name}_KiemTra{ext}"
        output_path = os.path.join(dir_name, output_name)
        
        self.doc.save(output_path)
        return output_path